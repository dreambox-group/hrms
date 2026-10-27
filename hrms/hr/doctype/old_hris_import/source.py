# Copyright (c) 2026, Dreambox and contributors
# For license information, please see license.txt
"""Reading the old AppSheet spreadsheets, leniently.

The sheets were edited by hand for years: dates are M/D/YYYY text or real dates, values vary in
spacing ("PartTime" / "Part Time"), headers have stray spaces. Values are normalised here; what
still can't be read is reported by the importer, never guessed silently.
"""

import re
from datetime import date, datetime

import openpyxl


class Workbook:
	"""One .xlsx, read once. Sheets are looked up by name, ignoring case and spacing."""

	def __init__(self, path):
		self.book = openpyxl.load_workbook(path, read_only=True, data_only=True)
		self.sheets = {key(name): name for name in self.book.sheetnames}

	def has(self, sheet):
		return key(sheet) in self.sheets

	def rows(self, sheet):
		"""Rows of one sheet as dicts keyed by trimmed header, with `_row` = the sheet's row number."""
		name = self.sheets.get(key(sheet))
		if not name:
			return []
		return read_rows(self.book[name])

	def all_rows(self, required_headers):
		"""(sheet name, rows) of every sheet whose header has all the given columns."""
		for name in self.book.sheetnames:
			rows = read_rows(self.book[name])
			if rows and all(h in rows[0] for h in required_headers):
				yield name, rows


def read_rows(worksheet):
	values = worksheet.iter_rows(values_only=True)
	header = [text(h) for h in next(values, ())]
	records = []
	for number, row in enumerate(values, start=2):
		if not any(v not in (None, "") for v in row):
			continue
		record = {h: v for h, v in zip(header, row) if h}
		record["_row"] = number
		records.append(record)
	return records


def text(value):
	if value is None:
		return ""
	if isinstance(value, float) and value.is_integer():
		value = int(value)
	return re.sub(r"\s+", " ", str(value)).strip()


def key(value):
	"""Comparable form of a value: 'Full Time - Permanent' == 'FullTime-Permanent'."""
	return re.sub(r"[^a-z0-9]", "", text(value).lower())


def to_date(value):
	"""Real dates, M/D/YYYY text (the sheets' en-US locale) or ISO text. None if unreadable."""
	if value in (None, ""):
		return None
	if isinstance(value, datetime):
		return value.date()
	if isinstance(value, date):
		return value
	value = text(value)
	for fmt in ("%m/%d/%Y", "%m/%d/%Y %H:%M:%S", "%m/%d/%Y %H:%M", "%Y-%m-%d", "%Y-%m-%d %H:%M:%S", "%d-%b-%Y"):
		try:
			return datetime.strptime(value, fmt).date()
		except ValueError:
			continue
	return None


def to_datetime(value, on_date=None):
	"""Datetimes, M/D/YYYY H:MM:SS text, or a bare time combined with `on_date`."""
	if value in (None, ""):
		return None
	if isinstance(value, datetime):
		# a time-only cell comes back as 1899-12-30 hh:mm
		return datetime.combine(on_date, value.time()) if on_date and value.year < 1901 else value
	if hasattr(value, "hour"):
		return datetime.combine(on_date, value) if on_date else None
	value = text(value)
	for fmt in ("%m/%d/%Y %H:%M:%S", "%m/%d/%Y %H:%M", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M"):
		try:
			return datetime.strptime(value, fmt)
		except ValueError:
			continue
	if on_date:
		for fmt in ("%H:%M:%S", "%H:%M", "%I:%M:%S %p", "%I:%M %p"):
			try:
				return datetime.combine(on_date, datetime.strptime(value, fmt).time())
			except ValueError:
				continue
	return None


def to_number(value, default=0.0):
	try:
		return float(text(value).replace(",", ".")) if text(value) else default
	except ValueError:
		return default


def to_latlong(value):
	"""AppSheet LatLong: '-6.1783, 106.6319'."""
	parts = re.findall(r"-?\d+(?:\.\d+)?", text(value))
	if len(parts) != 2:
		return None, None
	return float(parts[0]), float(parts[1])


def digits(value):
	return re.sub(r"\D", "", text(value))


def email(value):
	value = text(value).lower()
	return value if re.fullmatch(r"[^@\s]+@[^@\s]+\.[^@\s]+", value) else ""
